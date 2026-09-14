import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from pzrevize.database import Database
from pzrevize.external_influences import (
    DISPLAY_COLUMNS, apply_profile, blank_room, import_rooms_xlsx,
    parse_values_json, row_status, suggest_measures, classify_environment,
    abnormal_codes, minimum_ip, CHECKLIST_OPTIONS, CHECKLIST_MULTISELECT, checklist_requirement, abnormal_measure_entries, measure_entries, ABNORMAL_VALUES,
    parse_abnormal_measures_json, serialize_abnormal_measures, export_rooms_xlsx,
    GENERAL_PAGE_4_DEFAULT, GENERAL_PAGE_5_DEFAULT, REFERENCE_STANDARDS,
)


class ExternalInfluenceTests(unittest.TestCase):
    def test_mrazici_profile(self):
        room=apply_profile(blank_room(room_name='Mrazicí box'),'Mrazicí box')
        values=parse_values_json(room['values_json'])
        self.assertEqual(values['AA'],'AA8')
        self.assertEqual(values['AB'],'AB8')
        self.assertEqual(values['BA'],'BA4')
        self.assertEqual(values['BC'],'BC3')
        self.assertEqual(room['environment_class'],'ABNORMÁLNÍ')

    def test_conservative_measure_suggestions(self):
        values={c:'' for c in DISPLAY_COLUMNS}
        values.update({'AD':'AD4','BA':'BA5','BE':'BE2'})
        self.assertEqual(suggest_measures(values,'Rozvodna'),['7','8'])
        self.assertNotIn('10',suggest_measures(values,'Rozvodna'))
        self.assertIn('10',suggest_measures(values,'Nabíjení EV wallbox'))

    def test_xlsx_import_room_matrix(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'rooms.xlsx'
            wb=Workbook();ws=wb.active;ws.title='1.NP'
            ws.append(['','','','',''])
            ws.append(['Č.m.','Název místnosti']+DISPLAY_COLUMNS+['Opatření'])
            values=['AA5','AB5','AC1','AD1','AE1','AF1','AG1','AH1','AK1','AL1','AM-1-2','AN1','AP1','AQ2','AR1','---','BA1','BC2','BD1','BE1','CA1','CB1']
            ws.append(['1.01','Kancelář']+values+[''])
            wb.save(path)
            rooms=import_rooms_xlsx(path)
            self.assertEqual(len(rooms),1)
            self.assertEqual(rooms[0]['floor'],'1.NP')
            self.assertEqual(parse_values_json(rooms[0]['values_json'])['AM'],'AM-1-2')
            self.assertEqual(row_status(rooms[0]),'URČENO')


    def test_environment_classification_and_ip(self):
        values={c:'X' for c in DISPLAY_COLUMNS}
        values.update({
            'AA':'AA5','AB':'AB5','AC':'AC1','AD':'AD1','AE':'AE1','AF':'AF1',
            'AG':'AG1','AH':'AH1','AK':'AK1','AL':'AL1','AM':'AM-1-2','AN':'AN1',
            'AP':'AP1','AQ':'AQ2','AR':'AR1','AS':'AS1','BA':'BA1','BC':'BC2',
            'BD':'BD1','BE':'BE1','CA':'CA1','CB':'CB1',
        })
        self.assertEqual(classify_environment(values),'NORMÁLNÍ')
        values['AD']='AD4';values['BC']='BC3';values['AE']='AE3'
        self.assertEqual(classify_environment(values),'ABNORMÁLNÍ')
        self.assertIn('AD4',abnormal_codes(values))
        self.assertEqual(minimum_ip(values),('IPX4','IP4X'))

    def test_checklist_matches_uploaded_source_structure(self):
        self.assertEqual([x['code'] for x in CHECKLIST_OPTIONS['AD']], ['AD1','AD2','AD3','AD4','AD5','AD6','AD7','AD8'])
        self.assertEqual([x['code'] for x in CHECKLIST_OPTIONS['AL']], ['AL1','AL2'])
        self.assertEqual([x['code'] for x in CHECKLIST_OPTIONS['BE']], ['BE1','BE2','BE3','BE4'])
        self.assertEqual(CHECKLIST_MULTISELECT, {'AM','BA'})
        self.assertIn('IPX4',checklist_requirement('AD','AD4'))

    def test_audited_normal_classes_from_ed3_z1_z2(self):
        values={c:'X' for c in DISPLAY_COLUMNS}
        values.update({'AA':'AA4','AB':'AB4','AC':'AC1','AD':'AD1','AE':'AE1','AF':'AF1','AG':'AG1','AH':'AH1','AK':'AK1','AL':'AL1','AM':'AM-1-2','AN':'AN1','AP':'AP1','AQ':'AQ2','AR':'AR1','AS':'AS1','BA':'BA1','BC':'BC2','BD':'BD1','BE':'BE1','CA':'CA1','CB':'CB1'})
        self.assertEqual(classify_environment(values),'NORMÁLNÍ')
        self.assertNotIn('AA4',abnormal_codes(values))
        self.assertNotIn('AB4',abnormal_codes(values))
        self.assertNotIn('AQ2',abnormal_codes(values))
        values['BA']='BA4'
        self.assertEqual(classify_environment(values),'ABNORMÁLNÍ')
        self.assertIn('BA4',abnormal_codes(values))

    def test_legacy_codes_are_normalized_to_current_catalog(self):
        values=parse_values_json({'AM':'AM1-2','BE':'BE2N2'})
        self.assertEqual(values['AM'],'AM-1-2')
        self.assertEqual(values['BE'],'BE2')
        overrides=parse_abnormal_measures_json({'BE2N2':'Původní opatření','AM1-3':'EMC'})
        self.assertEqual(overrides['BE2'],'Původní opatření')
        self.assertEqual(overrides['AM-1-3'],'EMC')

    def test_database_has_room_matrix_columns(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/'test.db')
            cols={r[1] for r in db.fetchall('PRAGMA table_info(external_influences)')}
            for name in ('floor','room_no','room_name','values_json','measure_codes','environment_class','source_sheet','room_description','abnormal_measures_json','sort_order'):
                self.assertIn(name,cols)


    def test_reference_standards_and_formulas_are_seeded(self):
        codes=[x[0] for x in REFERENCE_STANDARDS]
        self.assertEqual(len(codes),13)
        self.assertIn('ČSN 33 2000-5-51 ed.3 + Z1, Z2',codes)
        self.assertIn('ČSN EN IEC 60079-10-1 ed.3',codes)
        self.assertIn('CLC/TR 50404',codes)
        self.assertIn('(dV/dt)min =',GENERAL_PAGE_4_DEFAULT)
        self.assertIn('Vz =',GENERAL_PAGE_5_DEFAULT)
        self.assertIn('t =',GENERAL_PAGE_5_DEFAULT)


    def test_every_abnormal_influence_has_traceable_measure(self):
        values={c:'' for c in DISPLAY_COLUMNS}
        values.update({'AD':'AD4','BC':'BC3','AR':'AR2','AN':'AN2'})
        entries=abnormal_measure_entries(values)
        self.assertEqual({e['code'] for e in entries},{'AD4','BC3','AR2','AN2'})
        self.assertTrue(all(e['requirement'].strip() for e in entries))
        self.assertTrue(all(e['complete'] for e in entries))


    def test_abnormal_measure_can_be_overridden_per_code(self):
        values={c:'' for c in DISPLAY_COLUMNS}
        values.update({'AD':'AD4','BC':'BC3'})
        overrides={'AD4':'Vlastní opatření pro vodu','BC3':'Vlastní ochranné opatření BC3'}
        raw=serialize_abnormal_measures(overrides)
        self.assertEqual(parse_abnormal_measures_json(raw),overrides)
        entries={e['code']:e for e in abnormal_measure_entries(values,raw)}
        self.assertEqual(entries['AD4']['requirement'],'Vlastní opatření pro vodu')
        self.assertEqual(entries['BC3']['requirement'],'Vlastní ochranné opatření BC3')
        self.assertTrue(entries['AD4']['custom'])

    def test_xlsx_export_contains_complete_matrix_and_measures(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'export.xlsx'
            room=blank_room(floor='1.NP',room_no='1.01',room_name='Technologie')
            values={c:'' for c in DISPLAY_COLUMNS}
            values.update({'AA':'AA5','AB':'AB5','AC':'AC1','AD':'AD4','AE':'AE1','AF':'AF1','AG':'AG1','AH':'AH1','AK':'AK1','AL':'AL1','AM':'AM-1-2','AN':'AN1','AP':'AP1','AQ':'AQ2','AR':'AR1','AS':'AS1','BA':'BA1','BC':'BC3','BD':'BD1','BE':'BE1','CA':'CA1','CB':'CB1'})
            room['values_json']=__import__('pzrevize.external_influences',fromlist=['serialize_values']).serialize_values(values)
            room['abnormal_measures_json']=serialize_abnormal_measures({'AD4':'IPX4 v určeném rozsahu','BC3':'Doplňková ochrana dle posouzení'})
            room['room_description']='Popis prostoru'
            export_rooms_xlsx(path,[room])
            from openpyxl import load_workbook
            ws=load_workbook(path,data_only=True).active
            headers=[c.value for c in ws[1]]
            for code in DISPLAY_COLUMNS:self.assertIn(code,headers)
            self.assertIn('Opatření / požadavky k VV',headers)
            self.assertIn('Popis prostoru',headers)
            row=[c.value for c in ws[2]]
            joined=' '.join(str(x or '') for x in row)
            self.assertIn('AD4: IPX4 v určeném rozsahu',joined)
            self.assertIn('BC3: Doplňková ochrana dle posouzení',joined)

    def test_ba4_is_abnormal_and_has_editable_measure(self):
        values={c:'' for c in DISPLAY_COLUMNS}
        values['BA']='BA4'
        entries={e['code']:e for e in measure_entries(values)}
        self.assertIn('BA4',entries)
        self.assertIn('poučen',entries['BA4']['requirement'].lower())
        self.assertIn('dohledem',entries['BA4']['requirement'].lower())
        self.assertTrue(entries['BA4']['abnormal'])

    def test_every_catalogued_abnormal_value_has_measure_template(self):
        for group,codes in ABNORMAL_VALUES.items():
            for code in codes:
                values={c:'' for c in DISPLAY_COLUMNS}; values[group]=code
                entries={e['code']:e for e in measure_entries(values)}
                self.assertIn(code,entries,(group,code))
                self.assertTrue(entries[code]['requirement'].strip(),(group,code))


    def test_2022_consolidated_normal_classes_and_current_codes(self):
        from pzrevize.external_influences import CHECKLIST_OPTIONS, classify_environment, BASE_HEATED
        base=dict(BASE_HEATED)
        for group,code in [('AA','AA4'),('AA','AA5'),('AB','AB4'),('AB','AB5'),('AQ','AQ2'),('BC','BC2')]:
            values=dict(base); values[group]=code
            self.assertEqual(classify_environment(values),'NORMÁLNÍ',(group,code))
        self.assertEqual([x['code'] for x in CHECKLIST_OPTIONS['BE']],['BE1','BE2','BE3','BE4'])
        self.assertIn('AM-1-2',[x['code'] for x in CHECKLIST_OPTIONS['AM']])
        self.assertNotIn('AM1-2',[x['code'] for x in CHECKLIST_OPTIONS['AM']])

    def test_2022_measure_texts_do_not_reuse_old_national_annex_details(self):
        from pzrevize.external_influences import checklist_item
        self.assertIn('speciálně navržené',checklist_item('AA1')['requirement'].lower())
        self.assertNotIn('ip20',checklist_item('AA1')['requirement'].lower())
        self.assertEqual(checklist_item('AD4')['source'],'tab. ZA.1, s. 21')
        self.assertIn('ipx4',checklist_item('AD4')['requirement'].lower())
        self.assertNotIn('ip44',checklist_item('AF2')['requirement'].lower())
        self.assertIn('dohledem',checklist_item('BA4')['requirement'].lower())

    def test_periodic_clone_keeps_external_room(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/'test.db')
            rid=db.execute("INSERT INTO revisions(revision_no,revision_type) VALUES(?,?)",('VV-1','VNEJSI'))
            room=apply_profile(blank_room(floor='1.NP',room_no='1.01',room_name='Kancelář'),'Běžná místnost – formát tabulky D1')
            db.execute("INSERT INTO external_influences(revision_id,floor,room_no,room_name,values_json,measure_codes,environment_class,result,sort_order) VALUES(?,?,?,?,?,?,?,?,?)",
                       (rid,room['floor'],room['room_no'],room['room_name'],room['values_json'],room['measure_codes'],room['environment_class'],'Určeno',1))
            new_id=db.clone_revision(rid,'periodic')
            row=db.fetchone('SELECT * FROM external_influences WHERE revision_id=?',(new_id,))
            self.assertEqual(row['room_no'],'1.01')
            self.assertEqual(parse_values_json(row['values_json'])['AA'],'AA5')
            self.assertEqual(row['result'],'K ověření')


if __name__=='__main__':
    unittest.main()
