from __future__ import annotations

import json
from datetime import datetime, date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


EXPECTED = {
    "NORMA": "standard",
    "NORMA - NÁZEV": "standard_name",
    "NORMA – NÁZEV": "standard_name",
    "ZÁVADY": "title",
    "ZAVADY": "title",
    "ČLÁNEK": "article",
    "CLANEK": "article",
    "POPIS": "requirement_text",
    "PLATNÁ": "valid_status",
    "PLATNA": "valid_status",
    "PLATNOST OD": "valid_from",
    "NAHRAZENA": "replaced_by",
    "KATEGORIE": "category",
    "KATEGORIE ZÁVADY": "defect_class",
    "KATEGORIE ZAVADY": "defect_class",
    "C1/C2/C3": "defect_class",
    "TEXT ZÁVADY": "defect_text",
    "TEXT ZAVADY": "defect_text",
    "ZÁVAŽNOST": "severity",
    "ZAVAZNOST": "severity",
    "POZNÁMKA": "note",
    "POZNAMKA": "note",
    "NORMOVÉ ODKAZY": "norm_refs_text",
    "NORMOVE ODKAZY": "norm_refs_text",
}


def _norm_header(v: Any) -> str:
    return " ".join(str(v or "").strip().upper().replace("_", " ").split())


def _date_to_text(v: Any) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        # Excel serial date. openpyxl normally converts if cell is formatted as date,
        # but older workbooks may leave it numeric.
        from datetime import timedelta
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(v))).date().isoformat()
        except Exception:
            return str(v)
    s = str(v).strip()
    return s


def _refs_from_text(text: Any) -> list[dict[str,str]]:
    refs=[]
    for raw in str(text or "").splitlines():
        line=raw.strip()
        if not line:continue
        parts=[x.strip() for x in line.split(" | ",2)]
        if len(parts)==1:parts=[parts[0],"",""]
        elif len(parts)==2:parts=[parts[0],parts[1],""]
        refs.append({"standard":parts[0],"article":parts[1],"citation":parts[2]})
    return refs


def _refs_to_text(raw: Any, standard: str="", article: str="", citation: str="") -> str:
    refs=[]
    try:
        data=json.loads(str(raw or "[]"))
        if isinstance(data,list):refs=[x for x in data if isinstance(x,dict)]
    except Exception:refs=[]
    if not refs and (standard or article or citation):refs=[{"standard":standard,"article":article,"citation":citation}]
    lines=[]
    for r in refs:
        st=str(r.get("standard") or "").strip();art=str(r.get("article") or "").strip();cit=str(r.get("citation") or "").strip()
        if st or art or cit:lines.append(f"{st} | {art} | {cit}".rstrip())
    return "\n".join(lines)


def inspect_headers(path: str | Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    return headers


def auto_mapping(headers: list[str]) -> dict[int, str]:
    result: dict[int, str] = {}
    for i, h in enumerate(headers):
        key = _norm_header(h)
        if key in EXPECTED:
            result[i] = EXPECTED[key]
    return result


def import_defects(path: str | Path, db, mapping: dict[int, str] | None = None, mode: str = "skip") -> dict[str, int]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [str(v or "").strip() for v in next(rows)]
    mapping = mapping or auto_mapping(headers)
    imported = skipped = updated = errors = 0

    for row in rows:
        try:
            item = {
                "category": "",
                "defect_class": "",
                "standard": "",
                "standard_name": "",
                "article": "",
                "title": "",
                "defect_text": "",
                "requirement_text": "",
                "severity": "",
                "valid_status": "",
                "valid_from": "",
                "valid_to": "",
                "replaced_by": "",
                "note": "",
                "norm_refs_text": "",
            }
            for idx, field in mapping.items():
                if idx >= len(row):
                    continue
                val = row[idx]
                if field in ("valid_from", "valid_to"):
                    item[field] = _date_to_text(val)
                else:
                    item[field] = "" if val is None else str(val).strip()

            if not any(item[k] for k in ("standard", "article", "title", "defect_text", "requirement_text")):
                continue
            if not item["defect_text"]:
                # dm-like catalog has a short category/title in ZÁVADY and the longer requirement in POPIS.
                # Preserve that information without pretending the POPIS is always the final defect wording.
                item["defect_text"] = item["title"] or item["requirement_text"]
            if not item["category"]:
                item["category"] = item["title"]
            if not item["valid_status"]:
                item["valid_status"] = "Neověřena"

            # 0.4.5: the actual severity is C1/C2/C3. Older spreadsheets may
            # contain that value in the former KATEGORIE ZÁVADY column.
            cls=(item.get("defect_class") or "").strip().upper()
            sev=(item.get("severity") or "").strip().upper()
            if cls in ("C1","C2","C3"):
                sev=cls
            item["severity"]=sev if sev in ("C1","C2","C3") else ""
            item["defect_class"]=item["severity"]
            refs=_refs_from_text(item.get("norm_refs_text"))
            if not refs and (item.get("standard") or item.get("article") or item.get("requirement_text")):
                refs=[{"standard":item.get("standard","").strip(),"article":item.get("article","").strip(),"citation":item.get("requirement_text","").strip()}]
            item["norm_refs_json"]=json.dumps(refs,ensure_ascii=False) if refs else ""
            if refs:
                item["standard"]=refs[0].get("standard","");item["article"]=refs[0].get("article","");item["requirement_text"]=refs[0].get("citation","")

            existing = db.fetchone(
                "SELECT id FROM defect_catalog WHERE TRIM(standard)=TRIM(?) AND TRIM(article)=TRIM(?) AND TRIM(defect_text)=TRIM(?) LIMIT 1",
                (item["standard"], item["article"], item["defect_text"]),
            )
            if existing:
                if mode == "update":
                    db.execute(
                        """UPDATE defect_catalog SET category=?,defect_class=?,standard_name=?,title=?,requirement_text=?,severity=?,valid_status=?,valid_from=?,valid_to=?,replaced_by=?,note=?,norm_refs_json=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                        (item["category"], item["defect_class"], item["standard_name"], item["title"], item["requirement_text"], item["severity"], item["valid_status"], item["valid_from"], item["valid_to"], item["replaced_by"], item["note"], item["norm_refs_json"], existing[0]),
                    )
                    updated += 1
                elif mode == "duplicate":
                    existing = None
                else:
                    skipped += 1
                    continue
            if not existing:
                db.execute(
                    """INSERT INTO defect_catalog(category,defect_class,standard,standard_name,article,norm_refs_json,title,defect_text,requirement_text,severity,valid_status,valid_from,valid_to,replaced_by,note)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (item["category"], item["defect_class"], item["standard"], item["standard_name"], item["article"], item["norm_refs_json"], item["title"], item["defect_text"], item["requirement_text"], item["severity"], item["valid_status"], item["valid_from"], item["valid_to"], item["replaced_by"], item["note"]),
                )
                imported += 1
        except Exception:
            errors += 1
    wb.close()
    return {"imported": imported, "updated": updated, "skipped": skipped, "errors": errors}


def export_defects(path: str | Path, db):
    rows = db.fetchall("""
        SELECT category,standard,standard_name,article,title,defect_text,requirement_text,severity,
               valid_status,valid_from,valid_to,replaced_by,note,norm_refs_json
        FROM defect_catalog ORDER BY standard,article,id
    """)
    wb = Workbook()
    ws = wb.active
    ws.title = "Závadovník"
    headers = [
        "KATEGORIE", "NORMA", "NORMA - NÁZEV", "ČLÁNEK", "ZÁVADY", "TEXT ZÁVADY",
        "POPIS", "ZÁVAŽNOST", "PLATNÁ", "PLATNOST OD", "PLATNOST DO", "NAHRAZENA", "POZNÁMKA", "NORMOVÉ ODKAZY"
    ]
    ws.append(headers)
    for r in rows:
        vals=list(r[:-1]);vals.append(_refs_to_text(r[-1],r[1] or "",r[3] or "",r[6] or ""));ws.append(vals)

    fill = PatternFill("solid", fgColor="202020")
    for c in ws[1]:
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = fill
        c.alignment = Alignment(vertical="center")
    widths = [18, 28, 36, 16, 30, 48, 52, 14, 14, 14, 14, 28, 32, 70]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)


def create_template(path: str | Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Závadovník"
    headers = [
        "KATEGORIE", "NORMA", "NORMA - NÁZEV", "ČLÁNEK", "ZÁVADY", "TEXT ZÁVADY",
        "POPIS", "ZÁVAŽNOST", "PLATNÁ", "PLATNOST OD", "PLATNOST DO", "NAHRAZENA", "POZNÁMKA", "NORMOVÉ ODKAZY"
    ]
    ws.append(headers)
    ws.append(["Ochrana před úrazem", "ČSN ...", "Název normy", "čl. ...", "Krátký název", "Formulace závady do revize", "Hlavní citace / podklad", "C2", "ANO", "", "", "", "", "ČSN ... | čl. ... | Citace / požadavek\nČSN ... | čl. ... | Další citace / požadavek"])
    for c in ws[1]:
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = PatternFill("solid", fgColor="202020")
    ws.freeze_panes = "A2"
    for i, w in enumerate([18,28,36,16,30,48,52,14,14,14,14,28,32,70], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)
